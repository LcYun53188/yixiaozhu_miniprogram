from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE


ROOT = Path(r"E:\Project\SauAPP")
SOURCE = ROOT / "docs" / "益小助-作品介绍文档.md"
OUTPUT = ROOT / "docs" / "益小助-作品介绍文档.docx"

BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK = RGBColor(0x1F, 0x4D, 0x78)
GRAY = RGBColor(0x66, 0x66, 0x66)


def set_font(run, size=11, bold=False, color=None, name="Calibri", east_asia="Microsoft YaHei"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def set_para(para, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=8, line=1.333):
    fmt = para.paragraph_format
    fmt.alignment = align
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    fmt.line_spacing = line


def add_hyperlink(paragraph, text, url, color="0563C1", underline=True):
    part = paragraph.part
    r_id = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")

    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    r_pr.append(c)

    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        r_pr.append(u)

    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Calibri")
    r_fonts.set(qn("w:hAnsi"), "Calibri")
    r_fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r_pr.append(r_fonts)

    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "22")
    r_pr.append(sz)

    new_run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    new_run.append(text_el)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def parse_meta(line):
    m = re.match(r"\*\*(.+?)\*\*：(.*)", line.strip())
    if not m:
        return None, None
    return m.group(1), m.group(2).strip()


def parse_markdown():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    return lines


def build_doc():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(1)
    sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)

    lines = parse_markdown()
    i = 0

    # Cover/title block
    while i < len(lines) and not lines[i].startswith("# "):
        i += 1
    title = lines[i][2:].strip() if i < len(lines) else "益小助"
    i += 1
    subtitle = lines[i][3:].strip() if i < len(lines) and lines[i].startswith("## ") else ""
    i += 1

    for _ in range(3):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    set_font(r, size=24, bold=True, color=DARK)
    set_para(p, align=WD_ALIGN_PARAGRAPH.CENTER, after=6, line=1.15)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(subtitle)
    set_font(r, size=14, color=GRAY)
    set_para(p, align=WD_ALIGN_PARAGRAPH.CENTER, after=18, line=1.15)

    while i < len(lines):
        line = lines[i].strip()
        if line == "---":
            i += 1
            break
        if not line:
            i += 1
            continue
        key, val = parse_meta(line)
        if key:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_para(p, align=WD_ALIGN_PARAGRAPH.CENTER, after=4, line=1.15)
            r = p.add_run(f"{key}：")
            set_font(r, size=11, bold=True)
            if key == "项目链接" and "](http" in val:
                m = re.match(r"\[(.+?)\]\((.+?)\)", val)
                if m:
                    add_hyperlink(p, m.group(1), m.group(2))
                else:
                    r2 = p.add_run(val)
                    set_font(r2, size=11)
            else:
                r2 = p.add_run(val)
                set_font(r2, size=11)
        i += 1

    doc.add_page_break()

    bullet_buffer = []
    number_buffer = []

    def flush_bullets():
        nonlocal bullet_buffer
        for item in bullet_buffer:
            p = doc.add_paragraph(style="List Bullet")
            set_para(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=4, line=1.2)
            r = p.add_run(item)
            set_font(r, size=11)
        bullet_buffer = []

    def flush_numbers():
        nonlocal number_buffer
        for item in number_buffer:
            p = doc.add_paragraph(style="List Number")
            set_para(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=4, line=1.2)
            r = p.add_run(item)
            set_font(r, size=11)
        number_buffer = []

    while i < len(lines):
        raw = lines[i]
        line = raw.strip()
        if not line:
            flush_bullets()
            flush_numbers()
            i += 1
            continue
        if line == "---":
            flush_bullets()
            flush_numbers()
            i += 1
            continue

        if line.startswith("## "):
            flush_bullets()
            flush_numbers()
            p = doc.add_paragraph()
            set_para(p, align=WD_ALIGN_PARAGRAPH.LEFT, before=18, after=10, line=1.15)
            r = p.add_run(line[3:].strip())
            set_font(r, size=16, bold=True, color=BLUE)
            i += 1
            continue

        if line.startswith("### "):
            flush_bullets()
            flush_numbers()
            p = doc.add_paragraph()
            set_para(p, align=WD_ALIGN_PARAGRAPH.LEFT, before=12, after=6, line=1.15)
            r = p.add_run(line[4:].strip())
            set_font(r, size=13, bold=True, color=BLUE)
            i += 1
            continue

        if line.startswith("#### "):
            flush_bullets()
            flush_numbers()
            p = doc.add_paragraph()
            set_para(p, align=WD_ALIGN_PARAGRAPH.LEFT, before=8, after=4, line=1.15)
            r = p.add_run(line[5:].strip())
            set_font(r, size=12, bold=True, color=DARK)
            i += 1
            continue

        m_num = re.match(r"^\d+\.\s+(.*)", line)
        if m_num:
            flush_bullets()
            number_buffer.append(m_num.group(1))
            i += 1
            continue

        if line.startswith("1. ") or line.startswith("2. "):
            flush_bullets()
            number_buffer.append(line.split(". ", 1)[1])
            i += 1
            continue

        if line.startswith("- "):
            flush_numbers()
            bullet_buffer.append(line[2:].strip())
            i += 1
            continue

        flush_bullets()
        flush_numbers()
        p = doc.add_paragraph()
        set_para(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=8, line=1.333)
        r = p.add_run(line)
        set_font(r, size=11)
        i += 1

    flush_bullets()
    flush_numbers()

    for sec in doc.sections:
        footer = sec.footer
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = p.add_run("益小助 作品介绍文档")
        set_font(r, size=9, color=GRAY)

    doc.save(OUTPUT)


if __name__ == "__main__":
    build_doc()
